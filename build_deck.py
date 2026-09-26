"""Build Token-Optimisation review deck (blue/white research theme)."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHT = RGBColor(0xDD, 0xEB, 0xF7)
DARK = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Emu(18288000)
prs.slide_height = Emu(10287000)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def bar(slide, title, kicker="TOKEN OPTIMISATION"):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    # top navy band
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Emu(1600000))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    tb = slide.shapes.add_textbox(Emu(700000), Emu(120000), W - Emu(1400000), Emu(500000))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = kicker
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(0xAE, 0xCC, 0xEA)
    p.font.bold = True
    tb2 = slide.shapes.add_textbox(Emu(700000), Emu(600000), W - Emu(1400000), Emu(800000))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(30)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    # footer
    ft = slide.shapes.add_textbox(Emu(700000), H - Emu(700000), W - Emu(2800000), Emu(400000))
    fp = ft.text_frame.paragraphs[0]
    fp.text = "Token Optimisation  •  Guide: Sonali Agrawal"
    fp.font.size = Pt(11)
    fp.font.color.rgb = GREY
    return Emu(2000000)


def body_box(slide, top, height=None):
    h = height or (H - top - Emu(1100000))
    tb = slide.shapes.add_textbox(Emu(900000), top, W - Emu(1800000), h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def bullets(tf, items, size=17, bold_first=False):
    for i, (head, tail) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.space_before = Pt(4)
        r = p.add_run()
        r.text = "▸  " + head
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if tail:
            r2 = p.add_run()
            r2.text = "  " + tail
            r2.font.size = Pt(size)
            r2.font.color.rgb = DARK


def cards(slide, top, items):
    n = len(items)
    gap, margin = Emu(400000), Emu(900000)
    cw = (W - 2 * margin - gap * (n - 1)) / n
    ch = Emu(3600000)
    for i, (head, tail) in enumerate(items):
        x = margin + i * (cw + gap)
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(x), top, int(cw), ch)
        shp.fill.solid()
        shp.fill.fore_color.rgb = LIGHT
        shp.line.color.rgb = BLUE
        shp.line.width = Pt(1.5)
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = head
        r.font.size = Pt(20)
        r.font.bold = True
        r.font.color.rgb = NAVY
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = tail
        r2.font.size = Pt(14)
        r2.font.color.rgb = DARK


# ---- 1 title ----
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
tb = s.shapes.add_textbox(Emu(900000), Emu(2200000), W - Emu(1800000), Emu(1200000))
p = tb.text_frame.paragraphs[0]
p.text = "TOKEN OPTIMISATION"
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER
tb2 = s.shapes.add_textbox(Emu(900000), Emu(3600000), W - Emu(1800000), Emu(1400000))
tb2.text_frame.word_wrap = True
p = tb2.text_frame.paragraphs[0]
p.text = "Cost-Efficient LLM Routing for Code-Mixed (Hinglish) Queries — compression, sweet-spot calibration and learned cascade routing under a statistical quality guarantee."
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(0xD9, 0xE2, 0xF3)
p.alignment = PP_ALIGN.CENTER
tb3 = s.shapes.add_textbox(Emu(900000), Emu(6200000), W - Emu(1800000), Emu(2000000))
tb3.text_frame.word_wrap = True
for i, line in enumerate(["Tushar Saharan  •  Shivam Kumar  •  Sahil Gupta", "Guide: Sonali Agrawal"]):
    p = tb3.text_frame.paragraphs[0] if i == 0 else tb3.text_frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = line
    p.font.size = Pt(20 if i == 0 else 17)
    p.font.bold = (i == 0)
    p.font.color.rgb = WHITE

# ---- 2 problem ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Research Problem")
bullets(body_box(s, top), [
    ("Tokenizer tax.", "English-biased tokenizers inflate Hinglish prompts 1.6–2.5x — users pay more per query."),
    ("Over-provisioning.", "Products route 100% of traffic to flagship models though 70–80% of routine queries suit cheap open models."),
    ("The gap.", "No calibrated, cost-aware router exists for code-mixed text: naive routing gambles quality, static rules waste money."),
])

# ---- 3 objectives ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Research Objectives")
cards(s, int(top) + Emu(400000), [
    ("~35%\nToken cut", "Lossless prompt compression with entity protection"),
    ("70%+\nCheap routing", "Cascade cheap 20B vs premium 120B by difficulty"),
    ("≤ 5%\nError bound", "Learn-Then-Test guarantee at 95% confidence"),
])

# ---- 4 background ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Background & Related Work")
bullets(body_box(s, top), [
    ("Cascade routing.", "FrugalGPT, RouteLLM, HybridLLM — learn when a small model suffices; mostly English, rarely cost-guaranteed."),
    ("Conformal control.", "Learn-Then-Test (Angelopoulos et al.) and Conformal Risk Control give finite-sample error budgets over thresholds."),
    ("Bandit routing.", "Contextual bandits directly optimise cost-vs-quality per query — the right formalism for single-step routing."),
    ("Missing piece.", "Nothing combines all three for Hinglish: script-mixed, entity-dense, price-sensitive traffic."),
])

# ---- 5 methodology ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Proposed Methodology")
steps = ["Compress\n35% fewer tokens", "Score\nd(x) difficulty", "Calibrate\nτ* sweet spot", "Route\ncheap / premium", "Receipt\naudit + saving"]
n = len(steps)
margin = Emu(700000)
bw = (W - 2 * margin) / n
y = int(top) + Emu(1400000)
bh = Emu(1800000)
for i, label in enumerate(steps):
    x = int(margin + i * bw + Emu(120000))
    w = int(bw - Emu(240000))
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, bh)
    shp.fill.solid()
    shp.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
    shp.line.color.rgb = BLUE
    shp.line.width = Pt(2)
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = label
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if i < n - 1:
        ax = x + w
        conn = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, ax, y + bh // 2 - Emu(150000), Emu(240000), Emu(300000))
        conn.fill.solid()
        conn.fill.fore_color.rgb = BLUE
        conn.line.fill.background()
note = s.shapes.add_textbox(Emu(900000), y + bh + Emu(500000), W - Emu(1800000), Emu(1200000))
note.text_frame.word_wrap = True
p = note.text_frame.paragraphs[0]
p.text = "d(x) = 0.40 code-mix + 0.30 entity + 0.20 math + 0.10 length  →  learned scorer  →  route cheap iff score < τ*."
p.font.size = Pt(15)
p.font.color.rgb = GREY
p.alignment = PP_ALIGN.CENTER

# ---- 6 compression (Tushar) ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Pillar 1 — Lossless Compression  (Tushar Saharan)")
bullets(body_box(s, top), [
    ("Entity-safe pruning.", "Regex masks emails, phones, amounts, math and code spans; compress around them; restore verbatim (fail-closed)."),
    ("Task reward.", "Reward = 0.7 answer-fidelity + 0.3 faithfulness — optimise correctness, not perplexity."),
    ("Controlled aggressiveness.", "Conformal Risk Control picks the max compression whose expected fidelity loss stays within budget α."),
])

# ---- 7 sweet spot (Shivam) ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Pillar 2 — The Sweet Spot  (Shivam Kumar)")
bullets(body_box(s, top), [
    ("α-sweep.", "Refit τ* across error budgets; cheap-share vs risk curve exposes diminishing returns."),
    ("Knee detection.", "Max-curvature point on the curve = the operating sweet spot: cheapest routing before risk climbs."),
    ("Mondrian calibration.", "Separate thresholds per code-mix bucket (low / mid / high) so heavy Hinglish isn't subsidised by easy English."),
])

# ---- 8 routing (Sahil) ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Pillar 3 — Learned Routing  (Sahil Gupta)")
bullets(body_box(s, top), [
    ("From heuristic to learned.", "MiniLM + difficulty features → P(fail | x) via LogReg / XGBoost; bandit head optimises cost − λ·failure directly."),
    ("Tiers.", "Cheap gpt-oss-20B vs premium-proxy gpt-oss-120B (free tier; flagship pricing in production design)."),
    ("Safe vs pure.", "Same policy, two thresholds: reward-optimal (pure RL) vs LTT-certified τ* (guaranteed) — the price of safety is measured, not assumed."),
])

# ---- 9 results ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Results  [DRAFT — refreshes from eval]")
rows = [
    ("Metric", "Heuristic", "Learned (safe)", "Pure RL"),
    ("Cheap-share @ α=0.05", "—", "—", "—"),
    ("Realized error [95% CI]", "—", "—", "—"),
    ("Saving vs all-premium", "—", "—", "—"),
    ("AUROC (fail prediction)", "—", "—", "—"),
]
tbl_shp = s.shapes.add_table(len(rows), 4, Emu(900000), int(top) + Emu(300000), W - Emu(1800000), Emu(4200000))
tbl = tbl_shp.table
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.text = val
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(15)
                r.font.bold = (i == 0)
                r.font.color.rgb = WHITE if i == 0 else (NAVY if j == 0 else DARK)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
note = s.shapes.add_textbox(Emu(900000), int(top) + Emu(4800000), W - Emu(1800000), Emu(800000))
note.text_frame.word_wrap = True
p = note.text_frame.paragraphs[0]
p.text = "n = 1200 Groq-labelled queries (794 / 203 / 203 splits), human-validated. CIs via 1000x bootstrap."
p.font.size = Pt(14)
p.font.color.rgb = GREY
p.alignment = PP_ALIGN.CENTER

# ---- 10 conclusion + team ----
s = prs.slides.add_slide(BLANK)
top = bar(s, "Conclusion, Limitations & Team")
bullets(body_box(s, top, Emu(5200000)), [
    ("Contribution.", "First calibrated cascade router for Hinglish: ~35% fewer tokens, majority-cheap routing, ≤5% error with proof."),
    ("Limits.", "Proxy tiers (not GPT-4o), 1200-query pilot, CPU-only training — all disclosed, all queued as future work."),
    ("Next.", "DistilBERT router, flagship premium, live pilot deployment."),
])
cards(s, H - Emu(2900000), [
    ("Tushar Saharan", "Compression"),
    ("Shivam Kumar", "Sweet-spot calibration"),
    ("Sahil Gupta", "Learned routing"),
])

prs.save("Token-Optimisation-Review.pptx")
print("saved Token-Optimisation-Review.pptx, slides:", len(prs.slides.__iter__.__self__._sldIdLst))
